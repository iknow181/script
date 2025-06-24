import os
import re
import time
import threading
import hashlib
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, scrolledtext, messagebox

import docx
import PyPDF2
import openpyxl
import xlrd
from PIL import Image

# 文本清洗与摘要
def clean_text(text, max_len=40):
    text = re.sub(r'[^\w\u4e00-\u9fff]+', '_', text).strip('_')
    return text[:max_len] or "Summary"

# 支持多编码读取 TXT
def read_txt(path):
    for enc in ['utf-8', 'gbk', 'utf-16', 'latin1']:
        try:
            with open(path, 'r', encoding=enc) as f:
                return clean_text(' '.join(f.readlines()[:3]))
        except:
            continue
    return None

# 内容摘要提取
def extract_summary(path: Path):
    suffix = path.suffix.lower()
    try:
        if suffix == '.pdf':
            with open(path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = ''.join((p.extract_text() or '') for p in reader.pages[:2])
                return clean_text(text)
        elif suffix == '.docx':
            return clean_text(' '.join(p.text for p in docx.Document(path).paragraphs[:3]))
        elif suffix == '.doc':
            return "DOC_File"
        elif suffix == '.txt':
            return read_txt(path)
        elif suffix == '.csv':
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                return clean_text(''.join(next(f) for _ in range(3)))
        elif suffix == '.xlsx':
            try:
                wb = openpyxl.load_workbook(path, read_only=True)
                ws = wb.active
                cells = [str(cell.value) for row in ws.iter_rows(max_row=3, max_col=5) for cell in row if
                         cell.value]
                wb.close()
                return clean_text(' '.join(cells))
            except Exception as e:
                print(f"[xlsx错误] {path}: {e}")
                return None
        elif suffix == '.xls':
            wb = xlrd.open_workbook(path)
            sheet = wb.sheet_by_index(0)
            cells = [str(sheet.cell_value(i, j)) for i in range(min(3, sheet.nrows)) for j in range(min(5, sheet.ncols))]
            return clean_text(' '.join(cells))
        elif suffix in ['.jpg', '.png']:
            with Image.open(path) as img:
                return f"{img.width}x{img.height}px"
    except Exception as e:
        return None
    return None

# 检查文件是否被占用
def is_file_locked(path):
    try:
        with open(path, 'rb'):
            return False
    except IOError:
        return True

# 重命名主逻辑
def rename_files_by_summary(root_dir, log_func):
    exts = ['.pdf', '.doc', '.docx', '.xlsx', '.xls', '.txt', '.csv', '.jpg', '.png']
    total, success, fail = 0, 0, 0

    for root, _, files in os.walk(root_dir):
        for fname in files:
            path = Path(root) / fname
            total += 1

            if path.suffix.lower() not in exts:
                continue

            if is_file_locked(path):
                log_func(f"[占用] {path.name} 被系统或程序占用，跳过")
                fail += 1
                continue

            try:
                summary = extract_summary(path)
                if not summary:
                    log_func(f"[跳过] {path.name}（内容为空或读取失败）")
                    fail += 1
                    continue

                # 哈希后缀保证唯一性
                suffix = path.suffix.lower()
                hash_tail = hashlib.md5(path.read_bytes()).hexdigest()[:8]
                base_name = f"{summary}_{hash_tail}"
                new_path = path.with_name(f"{base_name}{suffix}")

                # 冲突处理（最多 10 次）
                for i in range(10):
                    if not new_path.exists():
                        break
                    new_path = path.with_name(f"{base_name}_{i+1}{suffix}")
                else:
                    log_func(f"[跳过] {path.name}（重命名冲突过多）")
                    fail += 1
                    continue

                try:
                    path.rename(new_path)
                    log_func(f"[重命名] {path.name} → {new_path.name}")
                    success += 1
                except PermissionError:
                    log_func(f"[占用] {path.name} 首次失败，等待2秒重试...")
                    time.sleep(2)
                    try:
                        path.rename(new_path)
                        log_func(f"[重命名] {path.name} → {new_path.name}")
                        success += 1
                    except Exception as e2:
                        log_func(f"[异常] {path.name} 重命名失败：{e2}")
                        fail += 1

            except Exception as e:
                log_func(f"[异常] {path.name} 处理失败：{e}")
                fail += 1

    log_func(f"\n✅ 完成：共 {total} 个文件，成功 {success}，失败 {fail}\n")

# GUI 界面
def start_gui():
    def select_folder():
        path = filedialog.askdirectory()
        if path:
            folder_var.set(path)

    def log(msg):
        output.insert(tk.END, msg + "\n")
        output.yview_moveto(1)

    def run_thread():
        path = folder_var.get()
        if not os.path.isdir(path):
            messagebox.showerror("错误", "请选择一个有效的文件夹")
            return
        run_btn.config(state=tk.DISABLED)
        threading.Thread(target=lambda: [rename_files_by_summary(path, log), run_btn.config(state=tk.NORMAL)]).start()

    win = tk.Tk()
    win.title("SmartFileRenamer - 批量摘要重命名工具")
    win.geometry("750x520")

    tk.Label(win, text="选择待处理的文件夹:").pack(pady=5)
    frame = tk.Frame(win)
    frame.pack(pady=5)
    folder_var = tk.StringVar()
    tk.Entry(frame, textvariable=folder_var, width=65).pack(side=tk.LEFT, padx=5)
    tk.Button(frame, text="浏览", command=select_folder).pack(side=tk.LEFT)

    run_btn = tk.Button(win, text="开始重命名", command=run_thread, bg="#4CAF50", fg="white")
    run_btn.pack(pady=10)

    output = scrolledtext.ScrolledText(win, width=90, height=20)
    output.pack(pady=10)

    win.mainloop()

if __name__ == "__main__":
    start_gui()
