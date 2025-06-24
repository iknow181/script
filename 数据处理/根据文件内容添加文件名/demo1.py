import os
import re
from pathlib import Path

import docx
import PyPDF2
import openpyxl
import xlrd
import csv
from PIL import Image

# 清理提取的文本
def clean_text(text: str, max_len=30):
    cleaned = re.sub(r'[^\w\u4e00-\u9fff]+', '_', text).strip('_')
    return cleaned[:max_len] if cleaned else "摘要"

# 多编码尝试读取txt内容
def read_txt(path):
    for enc in ['utf-8', 'gbk', 'utf-16', 'latin1']:
        try:
            with open(path, 'r', encoding=enc) as f:
                lines = f.readlines()[:3]
                return clean_text(' '.join(line.strip() for line in lines))
        except Exception:
            continue
    return None

# 从各种文件中提取摘要
def extract_summary(path: Path):
    suffix = path.suffix.lower()
    try:
        if suffix == '.pdf':
            with open(path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = ""
                for page in reader.pages[:2]:
                    page_text = page.extract_text() or ''
                    text += page_text
                return clean_text(text)

        elif suffix == '.docx':
            doc = docx.Document(path)
            return clean_text(' '.join(p.text for p in doc.paragraphs[:3]))

        elif suffix == '.doc':
            return "[doc_文件]"

        elif suffix == '.txt':
            return read_txt(path)

        elif suffix == '.csv':
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                content = []
                for i, row in enumerate(reader):
                    if i >= 3: break
                    content.extend(row[:5])
                return clean_text(' '.join(content))

        elif suffix == '.xlsx':
            wb = openpyxl.load_workbook(path, read_only=True)
            ws = wb.active
            content = []
            for i, row in enumerate(ws.iter_rows(max_row=3, max_col=5)):
                for cell in row:
                    if cell.value:
                        content.append(str(cell.value))
            return clean_text(' '.join(content))

        elif suffix == '.xls':
            wb = xlrd.open_workbook(path)
            sheet = wb.sheet_by_index(0)
            content = []
            for row_idx in range(min(3, sheet.nrows)):
                row = sheet.row_values(row_idx)
                content.extend(row[:5])
            return clean_text(' '.join(str(x) for x in content))

        elif suffix in ['.png', '.jpg']:
            with Image.open(path) as img:
                return f"{img.width}x{img.height}px"

        else:
            return None

    except Exception as e:
        print(f"[读取失败] {path}：{e}")
        return None

# 主函数：重命名文件
def rename_files_by_summary(root_dir):
    exts = ['.pdf', '.doc', '.docx', '.xlsx', '.xls', '.txt', '.csv', '.jpg', '.png']
    for root, _, files in os.walk(root_dir):
        for fname in files:
            path = Path(root) / fname
            if path.suffix.lower() not in exts:
                continue
            summary = extract_summary(path)
            if not summary:
                continue

            new_name_base = summary
            new_path = path.with_name(f"{new_name_base}{path.suffix.lower()}")

            # 避免重名
            i = 1
            while new_path.exists():
                new_path = path.with_name(f"{new_name_base}_{i}{path.suffix.lower()}")
                i += 1

            print(f"[重命名] {path.name} → {new_path.name}")
            path.rename(new_path)

# 设置要处理的根目录
if __name__ == "__main__":
    folder_path = r"F:\workfile\20250613\test1"  # 修改为你自己的文件夹路径
    rename_files_by_summary(folder_path)
